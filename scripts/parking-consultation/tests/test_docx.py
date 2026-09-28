import io
from pathlib import Path

from docx import Document

from ses.docx_convert import convert_docx


def build_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("ΔΗΜΟΣ ΠΑΠΑΓΟΥ – ΧΟΛΑΡΓΟΥ")  # cover text, must be skipped
    document.add_heading("1\tΕΙΣΑΓΩΓΗ", level=1)
    document.add_paragraph("Εισαγωγικό κείμενο.")
    document.add_heading("1.1\tΑντικείμενο", level=2)
    paragraph = document.add_paragraph()
    paragraph.add_run("Επέκταση ζωνών ").bold = True
    paragraph.add_run("Α,Β").bold = True
    paragraph.add_run(" έως οδό Βενιζέλου.")
    document.add_paragraph("Πρώτο σημείο", style="List Bullet")
    document.add_paragraph("Δεύτερο σημείο", style="List Bullet")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "ΟΔΟΙ ΗΠΙΑΣ ΚΥΚΛΟΦΟΡΙΑΣ"
    table.cell(1, 0).text, table.cell(1, 1).text = "α/α", "ΟΔΟΣ"
    table.cell(2, 0).text, table.cell(2, 1).text = "1", "Ασπασίας (από Κύπρου έως Αναστάσεως)"
    document.add_paragraph("Πίνακας 1.1: Οδοί ήπιας κυκλοφορίας")
    document.add_heading("2\tΔΕΥΤΕΡΟ ΚΕΦΑΛΑΙΟ", level=1)
    document.add_heading("2.1\tΆρθρο", level=2)
    document.add_paragraph("Κείμενο άρθρου.")
    document.save(str(path))


def test_docx_becomes_chapters_articles_tables(tmp_path):
    source = tmp_path / "report.docx"
    build_docx(source)
    report = convert_docx(source, tmp_path / "figures", "/consultations/test")
    assert [c["num"] for c in report["chapters"]] == [1, 2]
    chapter = report["chapters"][0]
    assert chapter["title"] == "ΕΙΣΑΓΩΓΗ"
    assert chapter["preludeMd"].strip() == "Εισαγωγικό κείμενο."
    article = chapter["articles"][0]
    assert article["num"] == 1 and article["title"] == "Αντικείμενο"
    body = article["bodyMd"]
    assert "**Επέκταση ζωνών Α,Β** έως οδό Βενιζέλου." in body
    assert "- Πρώτο σημείο\n- Δεύτερο σημείο" in body
    assert "**ΟΔΟΙ ΗΠΙΑΣ ΚΥΚΛΟΦΟΡΙΑΣ**" in body
    assert "| α/α | ΟΔΟΣ |" in body and "| 1 | Ασπασίας (από Κύπρου έως Αναστάσεως) |" in body
    assert "*Πίνακας 1.1: Οδοί ήπιας κυκλοφορίας*" in body
    assert report["tables"][0]["title"] == "ΟΔΟΙ ΗΠΙΑΣ ΚΥΚΛΟΦΟΡΙΑΣ"
    assert report["tables"][0]["rows"] == [["1", "Ασπασίας (από Κύπρου έως Αναστάσεως)"]]
    assert report["anomalies"] == []
    assert report["source"]["sha256"]


def test_images_stay_in_their_section_and_the_last_one_is_kept(tmp_path):
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (4, 4), "red").save(buffer, format="PNG")
    png = buffer.getvalue()

    document = Document()
    document.add_heading("1\tΠΡΩΤΟ", level=1)
    document.add_heading("1.1\tΜε εικόνα", level=2)
    document.add_paragraph("Κείμενο.")
    document.add_picture(io.BytesIO(png))  # directly before the next heading, no caption
    document.add_heading("1.2\tΕπόμενο", level=2)
    document.add_paragraph("Άλλο κείμενο.")
    document.add_picture(io.BytesIO(png))  # the last thing in the document
    path = tmp_path / "report.docx"
    document.save(str(path))

    report = convert_docx(path, tmp_path / "figures", "/figures")
    first, second = report["chapters"][0]["articles"]
    assert first["bodyMd"].count("![") == 1
    assert second["bodyMd"].count("![") == 1
